"""Compare the tool's recalc results with Microsoft Excel's results, through COM automation.

Each group is a small workbook built with openpyxl. Its formulas are written the way Excel saves
them in a file (for example `_xlfn._xlws.FILTER`). Excel opens a copy of the workbook, each formula
is entered as Formula2 (the Microsoft 365 entry that spills arrays, as a user typing it would), and
Excel calculates. The tool's recalc output is then read back with openpyxl and compared with Excel:

- every formula value,
- every spill range (Excel's spill address against the array block the tool writes),
- every cell of every spill range,
- and the formula text Excel saves against the text the test wrote (checks the `_xlfn` prefixes).

The test never uses the tool to produce an expected value. Skipped when Excel or pywin32 is not
available, or when not running on Windows.
"""

import json
import math
import re
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "xlsx_tool.py"
EXCEL_EXE = Path(r"C:\Program Files\Microsoft Office\Root\Office16\EXCEL.EXE")
# Prefixes Excel writes in a file. Excel itself does not accept them when a formula is typed in.
FILE_PREFIX_RE = re.compile(r"_xlfn\._xlws\.|_xlfn\.|_xlpm\.")

# Shared data on sheet Data (columns A to H, rows 1 to 8), a horizontal table in K1:R3,
# and three defined names: ScoreCol, Keys (workbook scope) and Local (sheet scope on Calc).
DATA = {
    "A": [3, 8, 1, 8, 5, 2, 9, 4],
    "B": ["apple", "banana", "cherry", "banana", "date", "elder", "fig", "grape"],
    "C": ["x", "y", "x", "y", "x", "y", "x", "y"],
    "D": [10, 20, 30, 40, 50, 60, 70, 80],
    "E": [1, 2, 3, 4, 5, 8, 9, 12],
    "F": ["apple", "banana", "cherry", "date", "elder", "fig", "grape", "kiwi"],
    "G": [90, 70, 50, 40, 30, 20, 10, 5],
    "H": ["kiwi", "apple", "fig", "apple", "date", "kiwi", "banana", "cherry"],
}
HORIZONTAL = {
    "K": [1, "one", 100], "L": [2, "two", 200], "M": [3, "three", 300], "N": [4, "four", 400],
    "O": [5, "five", 500], "P": [6, "six", 600], "Q": [7, "seven", 700], "R": [8, "eight", 800],
}


def build_book(path, cells):
    """cells: list of (coordinate on sheet Calc, value or formula text)."""
    wb = openpyxl.Workbook()
    data = wb.active
    data.title = "Data"
    for col, values in DATA.items():
        for row, value in enumerate(values, start=1):
            data[f"{col}{row}"] = value
    for col, values in HORIZONTAL.items():
        for row, value in enumerate(values, start=1):
            data[f"{col}{row}"] = value
    calc = wb.create_sheet("Calc")
    for coord, value in cells:
        calc[coord] = value
    wb.defined_names["ScoreCol"] = DefinedName("ScoreCol", attr_text="Data!$D$1:$D$8")
    wb.defined_names["Keys"] = DefinedName("Keys", attr_text="Data!$E$1:$E$8")
    calc.defined_names["Local"] = DefinedName("Local", attr_text="Data!$E$1:$E$3")
    wb.save(path)


def layout(scalars=(), arrays=()):
    """Scalar formulas go down column A. Each array formula gets its own column group, row 1,
    eight columns apart, so spilled blocks never meet each other or the scalar column."""
    cells = [(f"A{i}", f) for i, f in enumerate(scalars, start=1)]
    cells += [(f"{get_column_letter(3 + 8 * k)}1", f) for k, f in enumerate(arrays)]
    return cells


# ---------------------------------------------------------------- the formula groups

SCALAR_MATH = [
    "=SUM(Data!A1:A8)", "=AVERAGE(Data!A1:A8)", "=MIN(Data!A1:A8)", "=MAX(Data!A1:A8)",
    "=COUNT(Data!A1:A8)", "=COUNTA(Data!B1:B8)", '=IF(SUM(Data!A1:A8)>30,"big","small")',
    '=IF(Data!A1>5,"hi","lo")', "=2^10", "=(Data!D1+Data!D2)*2-1", '="id-"&Data!A1', "=50%",
    "=3>2", "=Data!A1<>Data!A2", "=-Data!A1", "=SUM(Data!A1:A8*2)", "=SUM(ScoreCol)",
    "=SUM(Local)", "=SUM(Data!A1:A8)/COUNT(Data!A1:A8)", "=Data!A1&Data!B1",
    "=MAX(Data!A1:A8)-MIN(Data!A1:A8)", '=Data!B1="apple"', "=COUNTA(Data!A1:A8)",
    '=IF(Data!I1="","blank","filled")', "=Data!I1+5", "=Data!D8/Data!A1", "=1/3", "=2^-1",
]

CONDITIONAL = [
    '=SUMIF(Data!C1:C8,"x",Data!D1:D8)', '=SUMIF(Data!A1:A8,">4")',
    '=COUNTIF(Data!B1:B8,"banana")', '=COUNTIF(Data!A1:A8,">=5")', '=COUNTIF(Data!B1:B8,"b*")',
    '=SUMIFS(Data!D1:D8,Data!C1:C8,"y",Data!A1:A8,">2")',
    '=COUNTIFS(Data!C1:C8,"x",Data!A1:A8,">=3")',
    '=AVERAGEIFS(Data!D1:D8,Data!C1:C8,"x")',
    '=_xlfn.MAXIFS(Data!D1:D8,Data!C1:C8,"y")', '=_xlfn.MINIFS(Data!D1:D8,Data!C1:C8,"y")',
    '=_xlfn.TEXTJOIN(", ",TRUE,Data!B1:B3)', '=_xlfn.TEXTJOIN("-",FALSE,Data!A1:A3)',
    '=SUMIF(Data!B1:B8,"banana",Data!D1:D8)', "=COUNTIF(Data!A1:A8,8)",
    '=SUMIFS(ScoreCol,Data!C1:C8,"x")', '=COUNTIF(Data!H1:H8,"KIWI")',
    "=SUMIF(Data!A1:A8,8,Data!D1:D8)", '=AVERAGEIFS(Data!D1:D8,Data!A1:A8,">100")',
    '=_xlfn.MAXIFS(Data!D1:D8,Data!C1:C8,"zzz")', '=_xlfn.TEXTJOIN(";",TRUE,Data!H1:H3)',
]

EXACT_LOOKUPS = [
    "=VLOOKUP(8,Data!A1:D8,4,FALSE)", '=VLOOKUP("banana",Data!B1:D8,3,FALSE)',
    "=VLOOKUP(99,Data!A1:D8,2,FALSE)", "=HLOOKUP(5,Data!K1:R3,2,FALSE)",
    '=HLOOKUP("six",Data!K2:R3,2,FALSE)', "=INDEX(Data!D1:D8,3)", "=INDEX(Data!B1:C8,2,2)",
    "=INDEX(Data!A1:A8,MATCH(9,Data!A1:A8,0))", "=MATCH(8,Data!A1:A8,0)",
    '=MATCH("fig",Data!B1:B8,0)', '=INDEX(Data!D1:D8,MATCH("elder",Data!B1:B8,0))',
    '=IFERROR(VLOOKUP(99,Data!A1:D8,2,FALSE),"none")', '=VLOOKUP("APPLE",Data!B1:D8,2,FALSE)',
    "=MATCH(5,Data!A1:A8,0)", "=INDEX(Data!D1:D8,20)", '=MATCH("zzz",Data!B1:B8,0)',
    "=VLOOKUP(8,Data!A1:B8,2,FALSE)", "=HLOOKUP(7,Data!K1:R3,3,FALSE)",
    "=MATCH(8,Data!H1:H8,0)",
]

APPROXIMATE = [
    "=VLOOKUP(6,Data!E1:F8,2,TRUE)", "=VLOOKUP(100,Data!E1:F8,2,TRUE)",
    "=VLOOKUP(0,Data!E1:F8,2,TRUE)", "=HLOOKUP(6,Data!K1:R3,2,TRUE)",
    "=HLOOKUP(6,Data!K1:R3,3,TRUE)", "=MATCH(6,Data!E1:E8,1)", "=MATCH(6,Data!G1:G8,-1)",
    "=MATCH(4.5,Data!G1:G8,-1)", "=MATCH(0,Data!E1:E8,1)", "=MATCH(7,Data!E1:E8,1)",
    "=VLOOKUP(7,Data!E1:F8,2)", '=MATCH("c",Data!F1:F8,1)', '=VLOOKUP("dog",Data!F1:F8,1,TRUE)',
]

XLOOKUP_MODES = [
    "=_xlfn.XLOOKUP(8,Data!A1:A8,Data!D1:D8)",
    '=_xlfn.XLOOKUP(99,Data!A1:A8,Data!D1:D8,"none")',
    "=_xlfn.XLOOKUP(99,Data!A1:A8,Data!D1:D8)",
    "=_xlfn.XLOOKUP(8,Data!A1:A8,Data!D1:D8,,0,-1)",
    '=_xlfn.XLOOKUP("b*",Data!B1:B8,Data!D1:D8,"nf",2)',
    '=_xlfn.XLOOKUP("BANANA",Data!B1:B8,Data!D1:D8)',
    '=_xlfn.XLOOKUP(7,Data!A1:A8,Data!D1:D8,"nf",-1)',
    '=_xlfn.XLOOKUP(7,Data!A1:A8,Data!D1:D8,"nf",1)',
    '=_xlfn.XLOOKUP(7,Data!A1:A8,Data!D1:D8,"nf",1,-1)',
    '=_xlfn.XLOOKUP(6,Data!E1:E8,Data!F1:F8,"nf",-1,2)',
    '=_xlfn.XLOOKUP(6,Data!E1:E8,Data!F1:F8,"nf",1,2)',
    '=_xlfn.XLOOKUP(6,Data!G1:G8,Data!D1:D8,"nf",1,-2)',
    '=_xlfn.XLOOKUP(6,Data!G1:G8,Data!D1:D8,"nf",-1,-2)',
    '=_xlfn.XLOOKUP(12,Data!E1:E8,Data!E1:E8,"nf",0,2)',
    "=_xlfn.XLOOKUP(5,Data!A1:A8,Data!B1:B8)",
    "=_xlfn.XLOOKUP(3,Data!A1:A8,Data!A1:A8,\"nf\",0,1)",
    '=_xlfn.XLOOKUP("c*",Data!B1:B8,Data!B1:B8,"nf",2,-1)',
    '=_xlfn.XLOOKUP(4,Data!E1:E8,Data!E1:E8,"nf",-1,2)',
    '=_xlfn.XLOOKUP(10,Data!A1:A8,Data!D1:D8,"nf",-1,-1)',
]

DYNAMIC_ARRAYS = [
    "=_xlfn._xlws.FILTER(Data!A1:D8,Data!C1:C8=\"x\")",
    "=_xlfn._xlws.FILTER(Data!B1:B8,Data!A1:A8>4)",
    '=_xlfn._xlws.FILTER(Data!A1:A8,Data!A1:A8>100,"none")',
    "=_xlfn._xlws.FILTER(Data!A1:A8,Data!A1:A8>100)",
    '=_xlfn._xlws.FILTER(Data!A1:A8,Data!C1:C8="y")',
    "=_xlfn._xlws.FILTER(Data!E1:E8,Data!E1:E8)",
    "=_xlfn._xlws.SORT(Data!A1:A8)", "=_xlfn._xlws.SORT(Data!A1:A8,1,-1)",
    "=_xlfn._xlws.SORT(Data!A1:C8,3,1)", "=_xlfn._xlws.SORT(Data!A1:D8,4,-1)",
    "=_xlfn._xlws.SORT(Data!H1:H8)", "=_xlfn.UNIQUE(Data!H1:H8)", "=_xlfn.UNIQUE(Data!C1:C8)",
    "=_xlfn.UNIQUE(Data!A1:A8)", "=_xlfn.UNIQUE(Data!A1:A8,,TRUE)",
    "=_xlfn.UNIQUE(Data!C1:D8)", "=_xlfn.SEQUENCE(3,4)", "=_xlfn.SEQUENCE(4)",
    "=_xlfn.SEQUENCE(2,3,10,5)", "=_xlfn.SEQUENCE(0)", "=SUM(_xlfn.SEQUENCE(4))",
    "=_xlfn._xlws.SORT(ScoreCol,1,-1)", "=_xlfn._xlws.FILTER(Keys,Keys>3)",
    "=_xlfn.UNIQUE(Data!B1:B8,FALSE,TRUE)", "=_xlfn.XLOOKUP(8,Data!A1:A8,Data!B1:D8)",
    "=_xlfn._xlws.FILTER(Data!A1:D8,Data!A1:A8>4,\"none\")", "=_xlfn._xlws.SORT(Data!F1:F8,1,-1)",
    "=_xlfn._xlws.SORT(Data!B1:B8)",
    "=_xlfn.XLOOKUP(2,Data!K1:R1,Data!K1:R3)", '=_xlfn.XLOOKUP("seven",Data!K2:R2,Data!K3:R3)',
    "=_xlfn.XLOOKUP(7,Data!K1:R1,Data!K2:R2)",
]

LET_AND_NAMES = [
    "=_xlfn.LET(_xlpm.x,Data!A1,_xlpm.y,_xlpm.x*2,_xlpm.x+_xlpm.y)",
    "=_xlfn.LET(_xlpm.total,SUM(Data!A1:A8),_xlpm.total/COUNT(Data!A1:A8))",
    "=_xlfn.LET(_xlpm.r,_xlfn._xlws.FILTER(Data!A1:A8,Data!A1:A8>4),SUM(_xlpm.r))",
    "=_xlfn.LET(_xlpm.a,1,_xlpm.b,2,_xlpm.a+_xlpm.b)",
    "=_xlfn.LET(_xlpm.ScoreCol,5,_xlpm.ScoreCol+1)",
    "=_xlfn.LET(_xlpm.r,3,_xlpm.c,4,_xlpm.RC,5,_xlpm.r*_xlpm.c+_xlpm.RC)",
    '=_xlfn.LET(_xlpm.y,SUM(Data!A1:A8),IF(_xlpm.y>10,"big","small"))',
    "=SUM(Local)*2", "=SUM(Keys)",
]
LET_ARRAYS = [
    "=_xlfn.LET(_xlpm.x,_xlfn.SEQUENCE(3),_xlpm.x*10)",
    "=_xlfn.LET(_xlpm.n,Keys,_xlpm.n*2)",
]

ERRORS = [
    "=1/0", "=VLOOKUP(99,Data!A1:D8,2,FALSE)", "=Data!B1+1", "=Nosuchname+1", "=1E+307*10",
    "=INDEX(Data!D1:D8,20)", '=IFERROR(1/0,"fallback")', '=IFERROR(Data!B1+1,0)',
    "=_xlfn.XLOOKUP(99,Data!A1:A8,Data!D1:D8)", "=SUM(Data!B1:B8)", "=Data!A1/0",
    "=AVERAGE(Data!I1:I8)",
]

# Spill blocked by a constant (C3), and a spill blocked by a formula cell (E2) that sits inside
# the first formula's block. Excel shows #SPILL! in C1 and E1; the blocking cells keep their values.
SPILL_BLOCKING = [
    ("C1", "=_xlfn._xlws.SORT(Data!A1:A8)"), ("C3", "keep"),
    ("E1", "=_xlfn.SEQUENCE(3)"), ("E2", "=_xlfn.SEQUENCE(2)"),
]

GROUPS = {
    "scalar_math": lambda: layout(SCALAR_MATH),
    "conditional_aggregates": lambda: layout(CONDITIONAL),
    "exact_lookups": lambda: layout(EXACT_LOOKUPS),
    "approximate_lookups": lambda: layout(APPROXIMATE),
    "xlookup_modes": lambda: layout(XLOOKUP_MODES),
    "dynamic_arrays": lambda: layout(arrays=DYNAMIC_ARRAYS),
    "let_and_names": lambda: layout(LET_AND_NAMES, LET_ARRAYS),
    "errors": lambda: layout(ERRORS),
    "spill_blocking": lambda: list(SPILL_BLOCKING),
}

RESULTS = {}


# ---------------------------------------------------------------- Excel side

@pytest.fixture(scope="module")
def excel_app():
    if sys.platform != "win32":
        pytest.skip("Excel COM tests run on Windows only")
    if not EXCEL_EXE.exists():
        pytest.skip(f"Excel is not installed at {EXCEL_EXE}")
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        pytest.skip("pywin32 is not installed")
    pythoncom.CoInitialize()
    try:
        app = win32com.client.DispatchEx("Excel.Application")
    except Exception as exc:  # COM server missing, or Excel refuses to start
        pythoncom.CoUninitialize()
        pytest.skip(f"Excel did not start: {exc}")
    app.Visible = False
    app.DisplayAlerts = False
    try:
        yield app
    finally:
        app.Quit()
        pythoncom.CoUninitialize()


def read_used_values(ws):
    """Every cell of the used range on a sheet, as Excel shows it (errors as their text)."""
    values = {}
    used = ws.UsedRange
    top, left = used.Row, used.Column
    raw = used.Value2
    if not isinstance(raw, tuple):
        raw = ((raw,),)
    elif not isinstance(raw[0], tuple):
        raw = (raw,)
    for i, line in enumerate(raw):
        for j, value in enumerate(line):
            coord = f"{get_column_letter(left + j)}{top + i}"
            if isinstance(value, int) and not isinstance(value, bool):
                value = ws.Range(coord).Text  # an error value comes back as a number code
            values[coord] = value
    return values


def values_after_excel_recalc(app, path, work, anchors):
    """Open a file in Excel as it is (no formula is entered), calculate, and read the values and
    the spill ranges that Excel reports for the formula cells in anchors."""
    copy_path = work / ("excel_opened_" + path.name)
    shutil.copy(path, copy_path)
    wb = app.Workbooks.Open(str(copy_path), 0, True)
    try:
        app.CalculateFull()
        ws = wb.Worksheets("Calc")
        spills = {}
        for coord in anchors:
            rng = ws.Range(coord)
            if rng.HasSpill:
                spills[coord] = rng.SpillingToRange.Address.replace("$", "")
        return read_used_values(ws), spills
    finally:
        wb.Close(False)


def calculate_in_excel(app, source, work, formulas):
    """Open a copy of the workbook in Excel, enter each formula with Formula2, calculate, and
    return Excel's values, spill ranges and the formula text Excel saves."""
    copy_path = work / ("excel_" + source.name)
    saved_path = work / ("excel_saved_" + source.name)
    shutil.copy(source, copy_path)
    wb = app.Workbooks.Open(str(copy_path), 0, True)
    try:
        ws = wb.Worksheets("Calc")
        for coord, text in formulas:
            ws.Range(coord).Formula2 = FILE_PREFIX_RE.sub("", text)
        app.CalculateFull()

        values = read_used_values(ws)
        spills = {}
        for coord, _ in formulas:
            rng = ws.Range(coord)
            if rng.HasSpill:
                spills[coord] = rng.SpillingToRange.Address.replace("$", "")

        wb.SaveCopyAs(str(saved_path))
    finally:
        wb.Close(False)
    return values, spills, saved_path


# ---------------------------------------------------------------- comparison

def same_value(excel_value, ours):
    if excel_value is None:
        return ours is None
    if isinstance(excel_value, bool):
        return ours is excel_value
    if isinstance(excel_value, (int, float)):
        return (isinstance(ours, (int, float)) and not isinstance(ours, bool)
                and math.isclose(excel_value, ours, rel_tol=1e-9, abs_tol=1e-9))
    if isinstance(excel_value, str):
        return excel_value == ours
    return False


def cells_in(ref):
    """All cell coordinates in a range such as D1:F3 (or a single cell)."""
    from openpyxl.utils.cell import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    return [f"{get_column_letter(c)}{r}" for r in range(min_row, max_row + 1)
            for c in range(min_col, max_col + 1)]


def formula_text(value):
    return value.text if isinstance(value, ArrayFormula) else value


@pytest.mark.parametrize("group", sorted(GROUPS))
def test_recalc_matches_excel(excel_app, tmp_path, group):
    cells = GROUPS[group]()
    formulas = [(coord, value) for coord, value in cells if isinstance(value, str) and value.startswith("=")]
    source = tmp_path / f"{group}.xlsx"
    build_book(source, cells)

    proc = subprocess.run([sys.executable, str(SCRIPT), "recalc", str(source), "--json"],
                          capture_output=True, text=True, encoding="utf-8")
    assert proc.returncode == 0, proc.stdout + proc.stderr
    out = source.with_name(source.stem + ".recalc.xlsx")

    excel_values, excel_spills, saved_path = calculate_in_excel(excel_app, source, tmp_path, formulas)
    ours_values = openpyxl.load_workbook(out, data_only=True)["Calc"]
    ours_formulas = openpyxl.load_workbook(out)["Calc"]
    excel_saved = openpyxl.load_workbook(saved_path)["Calc"]

    anchors = {coord for coord, _ in formulas}
    spilled = set()
    for coord, addr in excel_spills.items():
        spilled.update(cell for cell in cells_in(addr) if cell != coord)

    formula_problems, spill_problems = [], []

    # 1. The formula text Excel saves must equal the text the test wrote (prefix check).
    for coord, text in formulas:
        saved = formula_text(excel_saved[coord].value)
        if saved != text:
            formula_problems.append(f"{coord}: Excel saves {saved!r}, input has {text!r}")

    # 2. Every formula and every spilled cell has the same value as Excel. Excel always gives a
    #    value for an entered formula, so an empty result for a formula is a failed check.
    empty = [coord for coord, _ in formulas if excel_values.get(coord) is None]
    assert not empty, f"Excel gave no value for {empty}"
    all_coords = set(excel_values) | {c.coordinate for row in ours_values.iter_rows() for c in row
                                      if c.value is not None}
    for coord in sorted(all_coords, key=lambda c: (int(re.sub(r"\D", "", c)),
                                                 openpyxl.utils.column_index_from_string(re.sub(r"\d", "", c)))):
        excel_value = excel_values.get(coord)
        ours = ours_values[coord].value
        if not same_value(excel_value, ours):
            message = f"{coord}: Excel {excel_value!r}, tool {ours!r}"
            (formula_problems if coord in anchors else spill_problems).append(message)

    # 3. The tool's own write command must store the same formula text that Excel saves.
    patch = [{"sheet": "Calc", "cell": coord, "formula": FILE_PREFIX_RE.sub("", text)}
             for coord, text in formulas]
    patch_path = tmp_path / f"{group}_patch.json"
    patch_path.write_text(json.dumps(patch), encoding="utf-8")
    written = tmp_path / f"{group}_written.xlsx"
    proc = subprocess.run([sys.executable, str(SCRIPT), "write", str(source), "--patch", str(patch_path),
                           "--out", str(written), "--json"], capture_output=True, text=True, encoding="utf-8")
    assert proc.returncode == 0, proc.stdout + proc.stderr
    written_calc = openpyxl.load_workbook(written)["Calc"]
    for coord, text in formulas:
        saved = formula_text(excel_saved[coord].value)
        stored = formula_text(written_calc[coord].value)
        if stored != saved:
            formula_problems.append(f"{coord}: write stores {stored!r}, Excel saves {saved!r}")

    # 4. Excel opens the file the tool wrote, without changes, and calculates the same values.
    opened, opened_spills = values_after_excel_recalc(excel_app, out, tmp_path, sorted(anchors))
    for coord in sorted(set(opened) | set(excel_values)):
        if not same_value(excel_values.get(coord), opened.get(coord)):
            message = (f"{coord}: Excel on the original {excel_values.get(coord)!r}, "
                       f"Excel on the tool's file {opened.get(coord)!r}")
            (formula_problems if coord in anchors else spill_problems).append(message)
    for coord in sorted(anchors):
        if opened_spills.get(coord) != excel_spills.get(coord):
            formula_problems.append(f"{coord}: Excel spills to {excel_spills.get(coord)} on the original, "
                                    f"{opened_spills.get(coord)} on the tool's file")

    # 5. Spill ranges: Excel's address against the array block the tool writes.
    for coord, _ in formulas:
        ref = ours_formulas[coord].value
        # a one-cell array formula (a formula that needs array evaluation) is not a spill
        ours_ref = ref.ref if isinstance(ref, ArrayFormula) and ref.ref != coord else None
        excel_addr = excel_spills.get(coord)
        if excel_addr is None and ours_ref is not None:
            formula_problems.append(f"{coord}: Excel does not spill, tool writes block {ours_ref}")
        elif excel_addr is not None and ours_ref != excel_addr:
            formula_problems.append(f"{coord}: Excel spills to {excel_addr}, tool writes {ours_ref}")

    RESULTS[group] = {
        "formulas": len(formulas),
        "formulas_differ": len(formula_problems),
        "spill_ranges": len(excel_spills),
        "spill_cells": len(spilled),
        "spill_cells_differ": len(spill_problems),
    }
    problems = formula_problems + spill_problems
    assert not problems, "\n".join(problems)


def test_summary_of_excel_comparison():
    if not RESULTS:
        pytest.skip("no Excel comparison ran")
    totals = {k: sum(r[k] for r in RESULTS.values()) for k in next(iter(RESULTS.values()))}
    print(f"\nExcel comparison: {totals}")
    assert totals["formulas_differ"] == 0 and totals["spill_cells_differ"] == 0
