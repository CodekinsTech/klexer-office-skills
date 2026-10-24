"""Hand-computed tests for the built-in formula evaluator.

Covers workbook and sheet defined names, approximate MATCH and VLOOKUP (binary search rules),
and XLOOKUP match and search modes. Every expected value is written out by hand in this file
(with the working in a comment), never taken from the evaluator. Refusals are checked for exit
code 1, the message and that no output file is kept.
"""

import json
import subprocess
import sys
from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "xlsx_tool.py"
NA = "#N/A"
REF = "#REF!"


def run(*args):
    return subprocess.run([sys.executable, str(SCRIPT), *[str(a) for a in args]],
                          capture_output=True, text=True, encoding="utf-8")


def make_book(path, columns, formulas, sheet="L"):
    """columns: {"A": [values from row 1 down]}; formulas: {"Q1": "=..."}. No cached values."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for col, values in columns.items():
        for row, value in enumerate(values, start=1):
            ws[f"{col}{row}"] = value
    for cell, formula in formulas.items():
        ws[cell] = formula
    wb.save(path)
    return path


def assert_values(book, expected, sheet="L"):
    proc = run("recalc", book, "--json")
    assert proc.returncode == 0, proc.stderr
    computed = json.loads(proc.stdout)["computed"]
    got = {cell: computed[f"{sheet}!{cell}"] for cell in expected}
    assert got == expected


def assert_refused(book, fragment):
    proc = run("recalc", book, "--json")
    assert proc.returncode == 1, proc.stdout + proc.stderr
    assert fragment in proc.stderr
    assert not book.with_name(book.stem + ".recalc.xlsx").exists()


# Shared data for the lookup tests, all on sheet L:
#   A1:A5 = 10, 20, 30, 40, 50             (ascending numbers)
#   B1:B5 = a, b, c, d, e                  (return values for A)
#   C1:C5 = 50, 40, 30, 20, 10             (descending numbers)
#   D1:D5 = v, w, x, y, z                  (return values for C)
#   G1:G4 = 10, 20, 20, 30                 (ascending, one repeated value in the middle)
#   H1:H3 = apple, apricot, banana         (text, for wildcards)
#   I1:I3 = 1, 2, 3                        (return values for H)
#   J1:J4 = 5, 7, 7, 9                     (ascending, a repeated 7)
#   K1:K4 = p, q, r, s                     (return values for J)
#   M1:O2 = 10 20 30 / x y z              (a two-row table for HLOOKUP)
BASE_COLUMNS = {
    "A": [10, 20, 30, 40, 50],
    "B": ["a", "b", "c", "d", "e"],
    "C": [50, 40, 30, 20, 10],
    "D": ["v", "w", "x", "y", "z"],
    "G": [10, 20, 20, 30],
    "H": ["apple", "apricot", "banana"],
    "I": [1, 2, 3],
    "J": [5, 7, 7, 9],
    "K": ["p", "q", "r", "s"],
    "M": [10, "x"],
    "N": [20, "y"],
    "O": [30, "z"],
}


def test_approximate_match_by_hand(tmp_path):
    formulas = {
        # MATCH type 1 (largest value <= key) on ascending A1:A5
        "Q1": "=MATCH(25,A1:A5,1)",    # 20 is the largest <= 25, at position 2
        "Q2": "=MATCH(50,A1:A5,1)",    # 50 is at position 5
        "Q3": "=MATCH(5,A1:A5,1)",     # nothing is <= 5, so #N/A
        "Q4": "=MATCH(99,A1:A5,1)",    # 50 is the largest <= 99, at position 5
        "Q5": "=MATCH(25,A1:A5)",      # match type omitted means 1, so position 2
        # MATCH type 1 on G1:G4 = 10, 20, 20, 30 (a repeat exists, but not of the chosen value)
        "Q6": "=MATCH(35,G1:G4,1)",    # 30 is the largest <= 35, unique, at position 4
        # MATCH type -1 (smallest value >= key) on descending C1:C5 = 50, 40, 30, 20, 10
        "Q7": "=MATCH(25,C1:C5,-1)",   # 30 is the smallest >= 25, at position 3
        "Q8": "=MATCH(10,C1:C5,-1)",   # 10 itself, at position 5
        "Q9": "=MATCH(60,C1:C5,-1)",   # no value >= 60, so #N/A
    }
    book = make_book(tmp_path / "approx.xlsx", BASE_COLUMNS, formulas)
    assert_values(book, {"Q1": 2, "Q2": 5, "Q3": NA, "Q4": 5, "Q5": 2,
                         "Q6": 4, "Q7": 3, "Q8": 5, "Q9": NA})


def test_approximate_vlookup_and_hlookup_by_hand(tmp_path):
    formulas = {
        # A1:B5 is sorted on column A; the largest A <= 35 is 30, row 3, so column B gives c
        "Q1": "=VLOOKUP(35,A1:B5,2,TRUE)",
        "Q2": "=VLOOKUP(35,A1:B5,2)",        # approximate by default
        "Q3": "=VLOOKUP(35,A1:B5,2,FALSE)",  # exact: 35 is not present, so #N/A
        "Q4": "=VLOOKUP(40,A1:B5,2,FALSE)",  # exact: 40 is in row 4, so d
        "Q5": "=VLOOKUP(5,A1:B5,2,TRUE)",    # nothing <= 5, so #N/A
        "Q6": "=VLOOKUP(10,A1:B5,3,TRUE)",   # column 3 does not exist in a 2-column range, so #REF!
        # M1:O2 is a sorted row 1 (10, 20, 30) over row 2 (x, y, z); largest <= 25 is 20, column 2
        "Q7": "=HLOOKUP(25,M1:O2,2,TRUE)",
    }
    book = make_book(tmp_path / "vh.xlsx", BASE_COLUMNS, formulas)
    assert_values(book, {"Q1": "c", "Q2": "c", "Q3": NA, "Q4": "d", "Q5": NA, "Q6": REF, "Q7": "y"})


def test_xlookup_match_modes_by_hand(tmp_path):
    formulas = {
        "Q1": '=XLOOKUP(30,A1:A5,B1:B5)',              # exact 30 is in row 3, so c
        "Q2": '=XLOOKUP(35,A1:A5,B1:B5,"none")',       # no exact 35, so the default none
        "Q3": '=XLOOKUP(35,A1:A5,B1:B5)',              # no default, so #N/A
        "Q4": '=XLOOKUP(35,A1:A5,B1:B5,"x",-1)',       # match -1: next smaller than 35 is 30, so c
        "Q5": '=XLOOKUP(35,A1:A5,B1:B5,"x",1)',        # match 1: next larger than 35 is 40, so d
        "Q6": '=XLOOKUP("ap*",H1:H3,I1:I3,"no",2)',    # match 2: apple matches ap*, first, so 1
        "Q7": '=XLOOKUP("b?n*",H1:H3,I1:I3,"no",2)',   # b, any one letter, n, then anything: banana, 3
        "Q8": '=XLOOKUP("c*",H1:H3,I1:I3,"no",2)',     # no text starts with c, so the default no
        "Q9": "=XLOOKUP(7,J1:J4,K1:K4)",                # search 1 (first to last): first 7 is row 2, q
        "Q10": '=XLOOKUP(7,J1:J4,K1:K4,"x",0,-1)',      # search -1 (last to first): last 7 is row 3, r
        "Q11": '=XLOOKUP(6,J1:J4,K1:K4,"x",-1)',        # next smaller than 6 is 5 (row 1), p
        "Q12": '=XLOOKUP(6,J1:J4,K1:K4,"x",1)',         # next larger than 6 is 7; first in search order is row 2, q
    }
    book = make_book(tmp_path / "xl_match.xlsx", BASE_COLUMNS, formulas)
    assert_values(book, {"Q1": "c", "Q2": "none", "Q3": NA, "Q4": "c", "Q5": "d", "Q6": 1,
                         "Q7": 3, "Q8": "no", "Q9": "q", "Q10": "r", "Q11": "p", "Q12": "q"})


def test_xlookup_binary_search_modes_by_hand(tmp_path):
    formulas = {
        # search 2 (binary, ascending A1:A5 = 10..50)
        "Q1": '=XLOOKUP(40,A1:A5,B1:B5,"x",0,2)',     # exact 40 is row 4, d
        "Q2": '=XLOOKUP(35,A1:A5,B1:B5,"x",0,2)',     # no exact 35, so the default x
        "Q3": '=XLOOKUP(35,A1:A5,B1:B5,"x",-1,2)',    # next smaller than 35 is 30, c
        "Q4": '=XLOOKUP(35,A1:A5,B1:B5,"x",1,2)',     # next larger than 35 is 40, d
        "Q5": '=XLOOKUP(5,A1:A5,B1:B5,"x",1,2)',      # next larger than 5 is 10, a
        "Q6": '=XLOOKUP(60,A1:A5,B1:B5,"x",-1,2)',    # next smaller than 60 is 50, e
        # search -2 (binary, descending C1:C5 = 50..10, return D1:D5 = v..z)
        "Q7": '=XLOOKUP(25,C1:C5,D1:D5,"x",-1,-2)',   # next smaller than 25 is 20, row 4, y
        "Q8": '=XLOOKUP(25,C1:C5,D1:D5,"x",1,-2)',    # next larger than 25 is 30, row 3, x
        "Q9": '=XLOOKUP(30,C1:C5,D1:D5,"x",0,-2)',    # exact 30 is row 3, x
        "Q10": '=XLOOKUP(50,C1:C5,D1:D5,"x",0,-2)',   # exact 50 is row 1, v
    }
    book = make_book(tmp_path / "xl_bin.xlsx", BASE_COLUMNS, formulas)
    assert_values(book, {"Q1": "d", "Q2": "x", "Q3": "c", "Q4": "d", "Q5": "a", "Q6": "e",
                         "Q7": "y", "Q8": "x", "Q9": "x", "Q10": "v"})


def test_unsorted_approximate_match_is_refused(tmp_path):
    columns = {"F": [30, 10, 20]}
    book = make_book(tmp_path / "unsorted.xlsx", columns, {"Q1": "=MATCH(15,F1:F3,1)"})
    assert_refused(book, "not sorted")


def test_approximate_match_refuses_repeated_chosen_value(tmp_path):
    # G1:G4 = 10, 20, 20, 30. The largest value <= 25 is 20, which is repeated, so refused.
    book = make_book(tmp_path / "dup.xlsx", {"G": [10, 20, 20, 30]}, {"Q1": "=MATCH(25,G1:G4,1)"})
    assert_refused(book, "occurs more than once")


def test_xlookup_binary_refusals(tmp_path):
    unsorted = make_book(tmp_path / "bin_unsorted.xlsx", {"F": [30, 10, 20], "K": ["p", "q", "r"]},
                         {"Q1": '=XLOOKUP(10,F1:F3,K1:K3,"x",0,2)'})
    assert_refused(unsorted, "not sorted")
    dup = make_book(tmp_path / "bin_dup.xlsx", {"G": [10, 20, 20, 30], "K": ["p", "q", "r", "s"]},
                    {"Q1": '=XLOOKUP(20,G1:G4,K1:K4,"x",0,2)'})
    assert_refused(dup, "occurs more than once")
    wild = make_book(tmp_path / "bin_wild.xlsx", BASE_COLUMNS, {"Q1": '=XLOOKUP("a*",H1:H3,I1:I3,"no",2,2)'})
    assert_refused(wild, "wildcard matching")


def test_match_type_other_than_three_is_refused(tmp_path):
    book = make_book(tmp_path / "mt.xlsx", BASE_COLUMNS, {"Q1": "=MATCH(10,A1:A5,2)"})
    assert_refused(book, "match_type 2 is not implemented")


def test_let_with_a_cell_reference_as_name_is_refused(tmp_path):
    # LET is implemented now (see test_dynamic_arrays.py); a name that looks like a cell is not valid.
    book = make_book(tmp_path / "let.xlsx", BASE_COLUMNS, {"Q1": "=LET(A1,1,A1)"})
    assert_refused(book, '"A1" is not a valid LET name')


def test_defined_names_by_hand(tmp_path):
    wb = openpyxl.Workbook()
    data = wb.active
    data.title = "Data"
    for row, value in enumerate([10, 20, 30], start=1):
        data.cell(row=row, column=1, value=value)
    data["B1"] = "=SUM(Rates)"          # 10 + 20 + 30 = 60
    data["B2"] = "=Top*2"               # Top is A3 = 30, so 60
    data["B3"] = "=sum(rates)+1"        # names are not case-sensitive: 60 + 1 = 61
    data["B4"] = "=Local+1"             # Local is A2 = 20, so 21
    data["B5"] = "=SUM(Other)"          # Notes!B1:B2 = 5 and 7, so 12
    data["B6"] = "=Missing+1"           # no such name anywhere in the workbook, so #NAME?
    data["B7"] = "=Quoted*5"            # 'My Data'!A1 = 4, so 20
    notes = wb.create_sheet("Notes")
    notes["B1"] = 5
    notes["B2"] = 7
    my = wb.create_sheet("My Data")
    my["A1"] = 4
    wb.defined_names["Rates"] = DefinedName("Rates", attr_text="Data!$A$1:$A$3")
    wb.defined_names["Top"] = DefinedName("Top", attr_text="Data!$A$3")
    wb.defined_names["Other"] = DefinedName("Other", attr_text="Notes!$B$1:$B$2")
    wb.defined_names["Quoted"] = DefinedName("Quoted", attr_text="'My Data'!$A$1")
    data.defined_names["Local"] = DefinedName("Local", attr_text="'Data'!$A$2")
    book = tmp_path / "names.xlsx"
    wb.save(book)
    assert_values(book, {"B1": 60, "B2": 60, "B3": 61, "B4": 21, "B5": 12, "B6": "#NAME?", "B7": 20},
                  sheet="Data")


def test_defined_name_to_constant_is_refused(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "L"
    ws["A1"] = 10
    ws["B1"] = "=TaxRate*A1"
    wb.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="0.2")
    book = tmp_path / "const_name.xlsx"
    wb.save(book)
    assert_refused(book, "does not refer to a single cell or range")
