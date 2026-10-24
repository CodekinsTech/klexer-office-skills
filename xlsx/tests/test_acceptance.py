"""Acceptance tests for xlsx_tool.py, one test per acceptance item in specs/xlsx.md section 3.

Input workbooks are built with openpyxl inside each test. Expected values are computed
here in plain Python, never with the tool's own evaluator.
"""

import datetime as dt
import json
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, Reference

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "xlsx_tool.py"
DATA_VALUES = [10, 20, 30]


def run(*args):
    return subprocess.run([sys.executable, str(SCRIPT), *[str(a) for a in args]],
                          capture_output=True, text=True, encoding="utf-8")


def make_data_book(path, extra=None):
    """Sheet Data: A1:A3 = 10, 20, 30; B1 = SUM(A1:A3) with no cached value. Sheet Notes is hidden."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for i, v in enumerate(DATA_VALUES, start=1):
        ws.cell(row=i, column=1, value=v)
    ws["B1"] = "=SUM(A1:A3)"
    ws["B2"] = "label"
    notes = wb.create_sheet("Notes")
    notes.sheet_state = "hidden"
    if extra:
        extra(ws)
    wb.save(path)
    return path


def read_json(proc):
    return json.loads(proc.stdout)


def test_sheets_lists_visible_and_hidden(tmp_path):
    book = make_data_book(tmp_path / "book.xlsx")
    proc = run("sheets", book, "--json")
    assert proc.returncode == 0, proc.stderr
    rows = {r["name"]: r for r in read_json(proc)["sheets"]}
    assert len(rows) == 2
    assert rows["Data"]["state"] == "visible"
    assert rows["Data"]["dimensions"] == "A1:B3"
    assert rows["Notes"]["state"] == "hidden"


def test_read_values_gives_null_and_not_calculated(tmp_path):
    book = make_data_book(tmp_path / "book.xlsx")
    proc = run("read", book, "--sheet", "Data", "--range", "A1:B1", "--json")
    assert proc.returncode == 0, proc.stderr
    out = read_json(proc)
    assert out["rows"] == [[DATA_VALUES[0], None]]
    assert out["not_calculated"] == ["B1"]


def test_read_formulas_returns_formula_text(tmp_path):
    book = make_data_book(tmp_path / "book.xlsx")
    proc = run("read", book, "--sheet", "Data", "--range", "A1:B1", "--mode", "formulas", "--json")
    assert proc.returncode == 0, proc.stderr
    assert read_json(proc)["rows"][0][1] == "=SUM(A1:A3)"


def test_write_formula_gets_xlfn_prefix(tmp_path):
    book = make_data_book(tmp_path / "book.xlsx")
    patch = tmp_path / "patch.json"
    patch.write_text(json.dumps([{"sheet": "Data", "cell": "C1",
                                  "formula": "=XLOOKUP(20,A1:A3,A1:A3)"}]), encoding="utf-8")
    proc = run("write", book, "--patch", patch, "--json")
    assert proc.returncode == 0, proc.stderr
    out = read_json(proc)
    assert {"cell": "Data!C1", "function": "XLOOKUP"} in out["prefixed_functions"]
    out_book = tmp_path / "book.write.xlsx"
    assert out["output"] == str(out_book)
    read_back = run("read", out_book, "--sheet", "Data", "--range", "C1", "--mode", "formulas", "--json")
    assert read_back.returncode == 0, read_back.stderr
    assert read_json(read_back)["rows"] == [["=_xlfn.XLOOKUP(20,A1:A3,A1:A3)"]]


def test_write_date_reads_back_iso(tmp_path):
    book = make_data_book(tmp_path / "book.xlsx")
    patch = tmp_path / "patch.json"
    patch.write_text(json.dumps([{"sheet": "Data", "cell": "D1", "value": "2026-10-24",
                                  "type": "date"}]), encoding="utf-8")
    proc = run("write", book, "--patch", patch, "--json")
    assert proc.returncode == 0, proc.stderr
    out_book = tmp_path / "book.write.xlsx"
    read_back = run("read", out_book, "--sheet", "Data", "--range", "D1", "--json")
    assert read_back.returncode == 0, read_back.stderr
    assert read_json(read_back)["rows"] == [["2026-10-24"]]


def test_format_currency_sets_number_format(tmp_path):
    book = make_data_book(tmp_path / "book.xlsx")
    proc = run("format", book, "--sheet", "Data", "--range", "A1:A3", "--preset", "currency", "--json")
    assert proc.returncode == 0, proc.stderr
    wb = openpyxl.load_workbook(tmp_path / "book.format.xlsx")
    formats = [wb["Data"].cell(row=r, column=1).number_format for r in range(1, 4)]
    assert formats == ['"$"#,##0.00'] * 3


def test_format_rejects_both_preset_and_code(tmp_path):
    book = make_data_book(tmp_path / "book.xlsx")
    proc = run("format", book, "--sheet", "Data", "--range", "A1:A3", "--preset", "currency", "--code", "0.0")
    assert proc.returncode == 1
    assert "Give exactly one of --preset or --code." in proc.stderr
    assert not (tmp_path / "book.format.xlsx").exists()


def test_recalc_stores_computed_values(tmp_path):
    book = make_data_book(tmp_path / "book.xlsx")
    proc = run("recalc", book, "--sheet", "Data", "--cells", "B1", "--json")
    assert proc.returncode == 0, proc.stderr
    out = read_json(proc)
    assert out["computed"] == {"Data!B1": sum(DATA_VALUES)}
    recalc_book = tmp_path / "book.recalc.xlsx"
    read_back = run("read", recalc_book, "--sheet", "Data", "--range", "A1:B1", "--json")
    assert read_back.returncode == 0, read_back.stderr
    result = read_json(read_back)
    assert result["rows"] == [[DATA_VALUES[0], sum(DATA_VALUES)]]
    assert result["not_calculated"] == []


def test_errors_reports_div_zero(tmp_path):
    def add_div_zero(ws):
        ws["C1"] = "=1/0"
    book = make_data_book(tmp_path / "book.xlsx", extra=add_div_zero)
    recalc = run("recalc", book, "--json")
    assert recalc.returncode == 0, recalc.stderr
    proc = run("errors", tmp_path / "book.recalc.xlsx", "--json")
    assert proc.returncode == 1
    out = read_json(proc)
    assert out["error_count"] == 1
    findings = [f for f in out["findings"] if f["value"] is not None]
    assert [(f["cell"], f["value"]) for f in findings] == [("C1", "#DIV/0!")]


def test_errors_clean_after_recalc(tmp_path):
    book = make_data_book(tmp_path / "book.xlsx")
    assert run("recalc", book, "--json").returncode == 0
    proc = run("errors", tmp_path / "book.recalc.xlsx", "--json")
    assert proc.returncode == 0, proc.stderr
    out = read_json(proc)
    assert out["error_count"] == 0
    assert out["not_calculated_count"] == 0


def test_validate_valid_workbook(tmp_path):
    book = make_data_book(tmp_path / "book.xlsx")
    proc = run("validate", book, "--json")
    assert proc.returncode == 0, proc.stderr
    assert read_json(proc)["verdict"] == "valid"


def test_validate_truncated_workbook_fails(tmp_path):
    book = make_data_book(tmp_path / "book.xlsx")
    broken = tmp_path / "broken.xlsx"
    shutil.copyfile(book, broken)
    with open(broken, "r+b") as fh:
        fh.truncate(1000)
    proc = run("validate", broken, "--json")
    assert proc.returncode == 1
    assert read_json(proc)["verdict"] == "invalid"


def _people_sheet(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "People"
    for col, title in enumerate(["Name", "Age", "Start", "Code"], start=1):
        ws.cell(row=1, column=col, value=title)
    for r, (name, age, start, code) in enumerate(rows, start=2):
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=age)
        cell = ws.cell(row=r, column=3, value=start)
        if isinstance(start, dt.date):
            cell.number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=4, value=code)
    wb.save(path)
    return path


RULES = {"columns": {
    "Name": {"type": "text", "required": True},
    "Age": {"type": "integer", "min": 0, "max": 120},
    "Start": {"type": "date"},
    "Code": {"type": "text", "pattern": "[A-Z][0-9]"},
}}


def test_check_values_reports_expected_violations(tmp_path):
    rows = [
        ("Ana", 34, dt.date(2026, 1, 5), "A1"),
        ("Ben", -2, "2026-02-30", "B2"),
        ("", 40, dt.date(2026, 3, 1), "C3"),
    ]
    book = _people_sheet(tmp_path / "people.xlsx", rows)
    rules = tmp_path / "rules.json"
    rules.write_text(json.dumps(RULES), encoding="utf-8")
    proc = run("check-values", book, "--rules", rules, "--sheet", "People", "--json")
    assert proc.returncode == 1
    out = read_json(proc)
    assert out["rows_checked"] == 3
    found = {(v["cell"], v["rule"]) for v in out["violations"]}
    # Spec expects B3 (Age below min), C3 (Start is text, not a real date) and A4 (Name empty).
    assert found == {("B3", "min"), ("C3", "type"), ("A4", "required")}


def test_check_values_clean_data_passes(tmp_path):
    rows = [
        ("Ana", 34, dt.date(2026, 1, 5), "A1"),
        ("Cara", 51, dt.date(2026, 3, 1), "C3"),
    ]
    book = _people_sheet(tmp_path / "people.xlsx", rows)
    rules = tmp_path / "rules.json"
    rules.write_text(json.dumps(RULES), encoding="utf-8")
    proc = run("check-values", book, "--rules", rules, "--sheet", "People", "--json")
    assert proc.returncode == 0, proc.stderr
    out = read_json(proc)
    assert out["violations"] == []
    assert out["rows_checked"] == 2


def test_write_refused_for_chart_without_flag(tmp_path):
    def add_chart(ws):
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=3), titles_from_data=False)
        ws.add_chart(chart, "E5")
    book = make_data_book(tmp_path / "book.xlsx", extra=add_chart)
    patch = tmp_path / "patch.json"
    patch.write_text(json.dumps([{"sheet": "Data", "cell": "C1", "value": 5}]), encoding="utf-8")
    proc = run("write", book, "--patch", patch)
    assert proc.returncode == 2
    assert "drawing" in proc.stderr.lower()
    assert not (tmp_path / "book.write.xlsx").exists()
    allowed = run("write", book, "--patch", patch, "--allow-drawing-rewrite", "--json")
    assert allowed.returncode == 0, allowed.stderr
    assert read_json(allowed)["drawings_rewritten"] is True
