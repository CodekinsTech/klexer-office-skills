#!/usr/bin/env python3
"""Generate the xlsx benchmark: input files for each task, plus tasks.json.

Run with the benchmark venv:
    python make_tasks.py

Writes:
    cases/<task_id>/   input files (workbooks built with openpyxl, plus patch and rules JSON)
    tasks.json         task definitions, including the expected results

The expected results are typed in below by hand, from specs/xlsx.md, SKILL.md and the
arithmetic noted next to each task. This script never runs xlsx_tool.py, so no expected
value comes from the tool's output.
"""

import datetime as dt
import json
import re
import shutil
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.defined_name import DefinedName

HERE = Path(__file__).resolve().parent
CASES = HERE / "cases"
TASKS_JSON = HERE / "tasks.json"

ERROR_CODES = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "#SPILL!", "#CALC!"}

TASKS = []


# ---------------------------------------------------------------- registration

def task(task_id, category, request, steps, build):
    TASKS.append({"id": task_id, "category": category, "request": request,
                  "steps": steps, "build": build})


def step(args, exit_code, checks=()):
    return {"args": list(args), "exit": exit_code, "checks": list(checks)}


# ---------------------------------------------------------------- check builders

def j_eq(path, equals, unordered=False):
    """Parsed stdout JSON at path equals the given value (dicts must match exactly)."""
    return {"check": "json", "path": path, "equals": equals, "unordered": unordered}


def j_contains(path, text):
    return {"check": "json_contains", "path": path, "text": text}


def j_truthy(path):
    return {"check": "json_truthy", "path": path}


def stderr_has(text):
    return {"check": "stderr_contains", "text": text}


def file_is(name, exists):
    return {"check": "file_exists", "file": name, "exists": exists}


def cell(file, sheet, ref, equals, mode="value"):
    """mode: value (cached values), formula (formula text), format (number_format)."""
    return {"check": "cell", "file": file, "sheet": sheet, "cell": ref, "mode": mode, "equals": equals}


# ---------------------------------------------------------------- file helpers

def new_book(first_sheet="Data"):
    wb = openpyxl.Workbook()
    wb.active.title = first_sheet
    return wb


def put(ws, cells):
    for ref, value in cells.items():
        ws[ref] = value


def write_json(path, obj):
    path.write_text(json.dumps(obj, indent=2, ensure_ascii=False), encoding="utf-8")


def store_cached(path, sheet_index, values):
    """Store cached results for formula cells the way Excel saves them.

    values: {"B1": 60} for a number, {"B1": "text"} for a string, {"B1": "#N/A"} for an error.
    The result is read back with openpyxl (data_only) and compared, so a bad patch fails here.
    """
    with zipfile.ZipFile(path) as zin:
        names = zin.namelist()
        parts = {n: zin.read(n) for n in names}
    part = f"xl/worksheets/sheet{sheet_index}.xml"
    xml = parts[part].decode("utf-8")
    for ref, val in values.items():
        pattern = re.compile(
            r'<c r="%s"([^>]*)>(<f[^>]*>.*?</f>|<f[^>]*/>)(?:<v\s*/>|<v>[^<]*</v>)?</c>' % ref, re.S)
        match = pattern.search(xml)
        if match is None:
            raise RuntimeError(f"formula cell {ref} not found in {part}")
        if isinstance(val, str) and val in ERROR_CODES:
            attr, text = ' t="e"', val
        elif isinstance(val, str):
            attr, text = ' t="str"', val.replace("&", "&amp;").replace("<", "&lt;")
        else:
            attr, text = "", repr(val)
        new = f'<c r="{ref}"{attr}>{match.group(2)}<v>{text}</v></c>'
        xml = xml[:match.start()] + new + xml[match.end():]
    parts[part] = xml.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, parts[n])
    check = openpyxl.load_workbook(path, data_only=True)
    ws = check.worksheets[sheet_index - 1]
    for ref, val in values.items():
        if ws[ref].value != val:
            raise RuntimeError(f"stored cache for {ref} reads back as {ws[ref].value!r}")


def make_macro_copy(src, dst):
    """Copy an xlsx and give it the macro-enabled content type, as Excel saves .xlsm."""
    with zipfile.ZipFile(src) as zin:
        items = [(n, zin.read(n)) for n in zin.namelist()]
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in items:
            if name == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml")
            zout.writestr(name, data)


# ---------------------------------------------------------------- builders

def build_sheets(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": 1, "B1": "x", "A2": 2, "B2": "y", "A3": 3, "B3": "z"})
    notes = wb.create_sheet("Notes")
    put(notes, {"A1": "note", "B1": "more", "A2": "a", "B2": "b"})
    notes.sheet_state = "hidden"
    wb.save(case / "book.xlsx")


def build_read_formula(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": 10, "A2": 20, "A3": 30, "B1": "=SUM(A1:A3)"})
    wb.save(case / "book.xlsx")


def build_read_cached(case):
    build_read_formula(case)
    store_cached(case / "book.xlsx", 1, {"B1": 60})


def build_refusals(case):
    (case / "legacy.xls").write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + bytes(504))
    (case / "data.csv").write_text("Name,Amount\nAna,10\nBen,20\n", encoding="utf-8")
    wb = new_book("Data")
    put(wb["Data"], {"A1": 10, "A2": 20})
    tmp = case / "_tmp.xlsx"
    wb.save(tmp)
    make_macro_copy(tmp, case / "macro.xlsm")
    tmp.unlink()


def build_write_source(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": 10, "A2": 20, "A3": 30})
    wb.save(case / "book.xlsx")


def build_write_date(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": "Invoice"})
    wb.save(case / "book.xlsx")
    write_json(case / "patch.json", [
        {"sheet": "Data", "cell": "D1", "value": "2026-10-24", "type": "date"}])


def build_write_value_types(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": 10, "A2": 20, "A3": 30, "B1": "label"})
    wb.save(case / "book.xlsx")
    write_json(case / "patch.json", [
        {"sheet": "Data", "cell": "A4", "value": 42},
        {"sheet": "Data", "cell": "B4", "value": "007"},
        {"sheet": "Data", "cell": "C4", "value": True},
        {"sheet": "Data", "cell": "A5", "value": "12.5", "type": "number"},
        {"sheet": "Data", "cell": "A3", "value": None},
    ])


def build_write_missing_equals(case):
    build_write_source(case)
    write_json(case / "patch.json", [
        {"sheet": "Data", "cell": "B1", "formula": "SUM(A1:A3)"}])


def build_write_bad_date(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": "x"})
    wb.save(case / "book.xlsx")
    write_json(case / "patch.json", [
        {"sheet": "Data", "cell": "D1", "value": "2026-02-30", "type": "date"}])


def build_write_xlookup(case):
    build_write_source(case)
    write_json(case / "patch.json", [
        {"sheet": "Data", "cell": "C1", "formula": "=XLOOKUP(20,A1:A3,A1:A3)"}])


def build_format(case):
    build_write_source(case)


def build_recalc_sum(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": 10, "A2": 20, "A3": 30,
                     "B1": "=SUM(A1:A3)", "B2": "=A1*2+A2/4", "B3": "=2^3"})
    wb.save(case / "book.xlsx")


def build_recalc_lookups(case):
    wb = new_book("Prices")
    put(wb["Prices"], {"A1": "Item", "B1": "Price",
                       "A2": "Apple", "B2": 1.5,
                       "A3": "Pear", "B3": 2,
                       "A4": "Plum", "B4": 3.25})
    orders = wb.create_sheet("Orders")
    put(orders, {"A1": "Pear", "A2": "Plum", "A3": "Apple",
                 "B1": "=VLOOKUP(A1,Prices!$A$2:$B$4,2,FALSE)",
                 "B2": "=INDEX(Prices!$B$2:$B$4,MATCH(A2,Prices!$A$2:$A$4,0))",
                 "B3": "=_xlfn.XLOOKUP(A3,Prices!$A$2:$A$4,Prices!$B$2:$B$4)",
                 "B4": '=_xlfn.XLOOKUP("Kiwi",Prices!$A$2:$A$4,Prices!$B$2:$B$4,"none")'})
    wb.save(case / "book.xlsx")


def build_recalc_sumifs(case):
    wb = new_book("Sales")
    rows = [("Region", "Rep", "Amount"),
            ("East", "Ana", 100),
            ("West", "Ben", 200),
            ("East", "Ben", 300),
            ("West", "Ana", 50),
            ("East", "Cara", 150)]
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            wb["Sales"].cell(row=r, column=c, value=value)
    put(wb["Sales"], {
        "E1": '=SUMIFS(C2:C6,A2:A6,"East")',
        "E2": '=COUNTIFS(A2:A6,"West",C2:C6,">=100")',
        "E3": '=AVERAGEIFS(C2:C6,B2:B6,"Ben")',
        "E4": '=COUNTIF(A2:A6,"East")',
        "E5": '=SUMIF(B2:B6,"Ana",C2:C6)'})
    wb.save(case / "book.xlsx")


def build_recalc_let(case):
    wb = new_book("Calc")
    put(wb["Calc"], {"A1": 5,
                     "B1": "=_xlfn.LET(_xlpm.x,A1,_xlpm.y,_xlpm.x*2,_xlpm.x+_xlpm.y)",
                     "B2": '=IFERROR(1/0,"n/a")',
                     "B3": '=IFERROR(A1*2,"n/a")'})
    wb.save(case / "book.xlsx")


def build_recalc_dynamic(case):
    wb = new_book("Staff")
    rows = [("Name", "Dept"), ("Cara", "Sales"), ("Ben", "Ops"), ("Ana", "Sales"),
            ("Dan", "Ops"), ("Eli", "Sales")]
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            wb["Staff"].cell(row=r, column=c, value=value)
    put(wb["Staff"], {
        "D1": '=_xlfn._xlws.FILTER(A2:A6,B2:B6="Sales")',
        "F1": "=_xlfn._xlws.SORT(A2:A6)",
        "H1": "=_xlfn.UNIQUE(B2:B6)",
        "J1": "=_xlfn.SEQUENCE(3,2)"})
    wb.save(case / "book.xlsx")


def build_recalc_defined_name(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": 10, "A2": 20, "A3": 30, "B1": 0.5,
                     "C1": "=A1*Rate", "C2": "=SUM(A1:A3)*Rate", "C3": "=Rate"})
    # Workbook-scope name for one cell (openpyxl 3.1 API)
    wb.defined_names["Rate"] = DefinedName("Rate", attr_text="Data!$B$1")
    wb.save(case / "book.xlsx")


def build_recalc_cross_sheet(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": 10, "A2": 20, "A3": 30, "A4": 40})
    put(wb.create_sheet("Summary"), {
        "B1": "=SUM(Data!A1:A4)",
        "B2": '=_xlfn.TEXTJOIN(", ",TRUE,Data!A1:A4)',
        "B3": '=_xlfn.MAXIFS(Data!A1:A4,Data!A1:A4,"<35")',
        "B4": '=_xlfn.MINIFS(Data!A1:A4,Data!A1:A4,">15")'})
    wb.save(case / "book.xlsx")


def build_recalc_spill_block(case):
    wb = new_book("Blk")
    put(wb["Blk"], {"A1": "=_xlfn.SEQUENCE(3)", "A2": "x"})
    wb.save(case / "book.xlsx")


def build_recalc_lambda(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": 10, "A2": 20, "A3": 30,
                     "B1": "=_xlfn.LAMBDA(_xlpm.x,_xlpm.x*2)(5)"})
    wb.save(case / "book.xlsx")


def build_recalc_unsorted(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": 30, "A2": 10, "A3": 20, "B1": "=VLOOKUP(15,A1:A3,1,TRUE)"})
    wb.save(case / "book.xlsx")


def build_errors_div0(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": 10, "A2": 20, "A3": 30, "B1": "=SUM(A1:A3)", "C1": "=1/0"})
    wb.save(case / "book.xlsx")


def build_errors_not_calculated(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": 10, "A2": 20, "A3": 30, "B1": "=SUM(A1:A3)"})
    wb.save(case / "book.xlsx")


def build_validate(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": 10, "A2": 20, "A3": 30, "B1": "=SUM(A1:A3)"})
    wb.save(case / "book.xlsx")
    data = (case / "book.xlsx").read_bytes()
    (case / "broken.xlsx").write_bytes(data[:1000])


def build_check_spec(case):
    wb = new_book("Data")
    ws = wb["Data"]
    rows = [("Name", "Age", "Start", "Code"),
            ("Ana", 34, dt.datetime(2026, 1, 5), "A1"),
            ("Ben", -2, "2026-02-30", "B2"),
            (None, 40, dt.datetime(2026, 3, 1), "C3")]
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    wb.save(case / "book.xlsx")
    write_json(case / "rules.json", {"columns": {
        "Name": {"type": "text", "required": True},
        "Age": {"type": "integer", "min": 0, "max": 120},
        "Start": {"type": "date"},
        "Code": {"type": "text", "pattern": "[A-Z][0-9]"}}})


def build_check_unique(case):
    wb = new_book("Data")
    ws = wb["Data"]
    rows = [("ID", "Status", "Qty", "Note"),
            ("X1", "open", 2.5, "ok"),
            ("X1", "pending", 11, "toolongtext"),
            ("X2", "closed", 5, "fine")]
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    wb.save(case / "book.xlsx")
    write_json(case / "rules.json", {"columns": {
        "ID": {"type": "text", "unique": True},
        "Status": {"type": "text", "allowed": ["open", "closed"]},
        "Qty": {"type": "decimal", "max": 10},
        "Note": {"type": "text", "max_length": 5}}})


def build_check_missing_column(case):
    wb = new_book("Data")
    put(wb["Data"], {"A1": "Name", "B1": "Age", "A2": "Ana", "B2": 34})
    wb.save(case / "book.xlsx")
    write_json(case / "rules.json", {"columns": {
        "Name": {"type": "text"},
        "Email": {"type": "text", "required": True}}})


def build_password(case):
    try:
        import msoffcrypto  # noqa: F401
    except ImportError as exc:
        raise ImportError("msoffcrypto-tool is not installed in the benchmark venv, "
                          "so no password-protected file can be made") from exc
    raise RuntimeError("password-protected case builder is not implemented for this msoffcrypto version")


def build_chart(case):
    wb = new_book("Data")
    ws = wb["Data"]
    put(ws, {"A1": 10, "A2": 20, "A3": 30})
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=3), titles_from_data=False)
    ws.add_chart(chart, "E2")
    wb.save(case / "chart.xlsx")
    write_json(case / "patch.json", [
        {"sheet": "Data", "cell": "D1", "formula": "=A1*2"}])


# ---------------------------------------------------------------- task definitions

# 1. sheets
task("sheets_lists_visible_and_hidden", "sheets",
     "What sheets are in book.xlsx, and which one is hidden?",
     [step(["sheets", "book.xlsx", "--json"], 0, [
         j_eq(["sheets"], [
             {"name": "Data", "state": "visible", "dimensions": "A1:B3", "max_row": 3, "max_column": 2},
             {"name": "Notes", "state": "hidden", "dimensions": "A1:B2", "max_row": 2, "max_column": 2}]),
     ])],
     build_sheets)

# 2-4. read
task("read_values_uncalculated_formula", "read",
     "Read A1:B1 of the Data sheet. The B1 formula has never been calculated.",
     [step(["read", "book.xlsx", "--sheet", "Data", "--range", "A1:B1", "--json"], 0, [
         j_eq(["rows"], [[10, None]]),
         j_eq(["not_calculated"], ["B1"]),
     ])],
     build_read_formula)

task("read_formulas_mode", "read",
     "Show me the formula text in B1 of the Data sheet, not the values.",
     [step(["read", "book.xlsx", "--sheet", "Data", "--range", "A1:B1", "--mode", "formulas", "--json"], 0, [
         j_eq(["rows"], [[10, "=SUM(A1:A3)"]]),
         j_eq(["not_calculated"], []),
     ])],
     build_read_formula)

task("read_cached_values_saved_file", "read",
     "The file was saved in Excel and has a stored total in B1. Read A1:B1.",
     [step(["read", "book.xlsx", "--sheet", "Data", "--range", "A1:B1", "--json"], 0, [
         j_eq(["rows"], [[10, 60]]),
         j_eq(["not_calculated"], []),
     ])],
     build_read_cached)

# 5. refusals of missing and unsupported files
task("refuse_missing_and_unsupported_files", "refusal",
     "Show the sheets of missing.xlsx, legacy.xls, data.csv and macro.xlsm, one at a time.",
     [step(["sheets", "missing.xlsx", "--json"], 1, [
         stderr_has("Input file not found: missing.xlsx")]),
      step(["sheets", "legacy.xls", "--json"], 2, [
         stderr_has("Legacy .xls is not supported. Save as .xlsx and try again.")]),
      step(["sheets", "data.csv", "--json"], 2, [
         stderr_has("CSV and TSV are not workbooks.")]),
      step(["sheets", "macro.xlsm", "--json"], 2, [
         stderr_has("Macro-enabled files are not supported. Save as .xlsx (macros would be dropped).")]),
      ],
     build_refusals)

# 6-10. write
task("write_xlookup_xlfn_prefix", "write",
     "Add an XLOOKUP in C1 that finds 20 in A1:A3 and returns the match. Save it as result.xlsx.",
     [step(["write", "book.xlsx", "--patch", "patch.json", "--out", "result.xlsx", "--json"], 0, [
         j_contains(["prefixed_functions"], "XLOOKUP"),
         file_is("result.xlsx", True),
         cell("result.xlsx", "Data", "C1", "=_xlfn.XLOOKUP(20,A1:A3,A1:A3)", mode="formula"),
     ])],
     build_write_xlookup)

task("write_date_cell", "write",
     "Put the date 2026-10-24 in D1 as a real date, not as text.",
     [step(["write", "book.xlsx", "--patch", "patch.json", "--out", "result.xlsx", "--json"], 0, [
         cell("result.xlsx", "Data", "D1", "2026-10-24", mode="value"),
         cell("result.xlsx", "Data", "D1", "yyyy-mm-dd", mode="format"),
     ])],
     build_write_date)

task("write_value_types_and_clear", "write",
     "Write 42 in A4, the text 007 in B4, TRUE in C4, the text 12.5 as a number in A5, and clear A3.",
     [step(["write", "book.xlsx", "--patch", "patch.json", "--out", "result.xlsx", "--json"], 0, [
         cell("result.xlsx", "Data", "A4", 42),
         cell("result.xlsx", "Data", "B4", "007"),
         cell("result.xlsx", "Data", "C4", True),
         cell("result.xlsx", "Data", "A5", 12.5),
         cell("result.xlsx", "Data", "A3", None),
     ])],
     build_write_value_types)

task("write_formula_missing_equals", "write",
     "Put SUM(A1:A3) in B1 as a formula.",
     [step(["write", "book.xlsx", "--patch", "patch.json", "--out", "result.xlsx", "--json"], 1, [
         stderr_has('Formula in B1 must start with "=".'),
         file_is("result.xlsx", False),
     ])],
     build_write_missing_equals)

task("write_invalid_date", "write",
     "Put 2026-02-30 in D1 as a date.",
     [step(["write", "book.xlsx", "--patch", "patch.json", "--out", "result.xlsx", "--json"], 1, [
         stderr_has('"2026-02-30" is not a valid YYYY-MM-DD date (cell D1).'),
         file_is("result.xlsx", False),
     ])],
     build_write_bad_date)

# 11-12. format
task("format_currency_preset", "format",
     "Format A1:A3 on the Data sheet as currency.",
     [step(["format", "book.xlsx", "--sheet", "Data", "--range", "A1:A3", "--preset", "currency",
            "--out", "result.xlsx", "--json"], 0, [
         cell("result.xlsx", "Data", "A1", '"$"#,##0.00', mode="format"),
         cell("result.xlsx", "Data", "A2", '"$"#,##0.00', mode="format"),
         cell("result.xlsx", "Data", "A3", '"$"#,##0.00', mode="format"),
         cell("result.xlsx", "Data", "A2", 20),
     ])],
     build_format)

task("format_both_preset_and_code_refused", "format",
     "Format A1:A3 as currency with the code 0.0 at the same time.",
     [step(["format", "book.xlsx", "--sheet", "Data", "--range", "A1:A3", "--preset", "currency",
            "--code", "0.0", "--out", "result.xlsx", "--json"], 1, [
         stderr_has("Give exactly one of --preset or --code."),
         file_is("result.xlsx", False),
     ])],
     build_format)

# 13-20. recalc and formula support
task("recalc_sum_and_arithmetic", "recalc",
     "Calculate B1 (sum of A1:A3), B2 (A1*2+A2/4) and B3 (2^3).",
     [step(["recalc", "book.xlsx", "--sheet", "Data", "--cells", "B1,B2,B3", "--out", "result.xlsx", "--json"], 0, [
         j_eq(["computed"], {"Data!B1": 60, "Data!B2": 25, "Data!B3": 8}),
         cell("result.xlsx", "Data", "B1", 60),
         cell("result.xlsx", "Data", "B2", 25),
         cell("result.xlsx", "Data", "B3", 8),
         cell("result.xlsx", "Data", "B1", "=SUM(A1:A3)", mode="formula"),
     ])],
     build_recalc_sum)

task("recalc_lookup_functions", "recalc",
     "On Orders, look up the price of each item in Prices with VLOOKUP, INDEX/MATCH and XLOOKUP. Kiwi is not in the list, so show 'none' for it.",
     [step(["recalc", "book.xlsx", "--sheet", "Orders", "--cells", "B1:B4", "--out", "result.xlsx", "--json"], 0, [
         j_eq(["computed"], {"Orders!B1": 2, "Orders!B2": 3.25, "Orders!B3": 1.5, "Orders!B4": "none"}),
         cell("result.xlsx", "Orders", "B2", 3.25),
     ])],
     build_recalc_lookups)

task("recalc_sumifs_family", "recalc",
     "For the Sales sheet: total East amount, count of West sales of 100 or more, average Ben amount, count of East rows, total for Ana.",
     [step(["recalc", "book.xlsx", "--sheet", "Sales", "--cells", "E1:E5", "--out", "result.xlsx", "--json"], 0, [
         j_eq(["computed"], {"Sales!E1": 550, "Sales!E2": 1, "Sales!E3": 250, "Sales!E4": 3, "Sales!E5": 150}),
         cell("result.xlsx", "Sales", "E3", 250),
     ])],
     build_recalc_sumifs)

task("recalc_let_and_iferror", "recalc",
     "With A1 = 5, B1 should be x + 2x with LET, and B2 and B3 should show n/a when there is an error.",
     [step(["recalc", "book.xlsx", "--sheet", "Calc", "--cells", "B1:B3", "--out", "result.xlsx", "--json"], 0, [
         j_eq(["computed"], {"Calc!B1": 15, "Calc!B2": "n/a", "Calc!B3": 10}),
         j_eq(["errors"], {}),
     ])],
     build_recalc_let)

task("recalc_filter_sort_unique_sequence", "recalc",
     "From the Staff list: FILTER the Sales names into D1, SORT all names into F1, UNIQUE departments into H1, and a 3x2 SEQUENCE into J1.",
     [step(["recalc", "book.xlsx", "--sheet", "Staff", "--cells", "D1,F1,H1,J1", "--out", "result.xlsx", "--json"], 0, [
         j_eq(["errors"], {}),
         cell("result.xlsx", "Staff", "D1", "Cara"),
         cell("result.xlsx", "Staff", "D2", "Ana"),
         cell("result.xlsx", "Staff", "D3", "Eli"),
         cell("result.xlsx", "Staff", "D4", None),
         cell("result.xlsx", "Staff", "F1", "Ana"),
         cell("result.xlsx", "Staff", "F5", "Eli"),
         cell("result.xlsx", "Staff", "F6", None),
         cell("result.xlsx", "Staff", "H1", "Sales"),
         cell("result.xlsx", "Staff", "H2", "Ops"),
         cell("result.xlsx", "Staff", "J1", 1),
         cell("result.xlsx", "Staff", "K1", 2),
         cell("result.xlsx", "Staff", "J2", 3),
         cell("result.xlsx", "Staff", "K2", 4),
         cell("result.xlsx", "Staff", "J3", 5),
         cell("result.xlsx", "Staff", "K3", 6),
     ])],
     build_recalc_dynamic)

task("recalc_defined_name", "recalc",
     "There is a workbook name Rate that points at B1 (0.5). Multiply A1, the sum of A1:A3, and Rate itself by it.",
     [step(["recalc", "book.xlsx", "--sheet", "Data", "--cells", "C1,C2,C3", "--out", "result.xlsx", "--json"], 0, [
         j_eq(["computed"], {"Data!C1": 5, "Data!C2": 30, "Data!C3": 0.5}),
         cell("result.xlsx", "Data", "C1", "=A1*Rate", mode="formula"),
     ])],
     build_recalc_defined_name)

task("recalc_cross_sheet_textjoin_maxifs", "recalc",
     "On Summary, total the Data column, join its values with commas, and show the largest value under 35 and the smallest over 15.",
     [step(["recalc", "book.xlsx", "--sheet", "Summary", "--cells", "B1:B4", "--out", "result.xlsx", "--json"], 0, [
         j_eq(["computed"], {"Summary!B1": 100, "Summary!B2": "10, 20, 30, 40", "Summary!B3": 30, "Summary!B4": 20}),
     ])],
     build_recalc_cross_sheet)

task("recalc_spill_blocked_shows_spill_error", "recalc",
     "A1 has SEQUENCE(3), but A2 already has text. The spill should be blocked and reported.",
     [step(["recalc", "book.xlsx", "--sheet", "Blk", "--cells", "A1", "--out", "result.xlsx", "--json"], 0, [
         j_eq(["errors"], {"Blk!A1": "#SPILL!"}),
         cell("result.xlsx", "Blk", "A2", "x"),
         cell("result.xlsx", "Blk", "A3", None),
     ])],
     build_recalc_spill_block)

# 21-22. refusals of unsupported formulas
task("refuse_lambda_on_recalc", "refusal",
     "Calculate B1, which uses LAMBDA(x, x*2)(5).",
     [step(["recalc", "book.xlsx", "--sheet", "Data", "--out", "result.xlsx", "--json"], 1, [
         stderr_has("Formula in Data!B1 is not supported by the built-in evaluator: "
                    "the function LAMBDA is not supported. Nothing was written."),
         file_is("result.xlsx", False),
     ])],
     build_recalc_lambda)

task("refuse_unsorted_approximate_lookup", "refusal",
     "Use VLOOKUP with approximate match on A1:A3, which is not sorted, and calculate it.",
     [step(["recalc", "book.xlsx", "--sheet", "Data", "--out", "result.xlsx", "--json"], 1, [
         stderr_has("Formula in Data!B1 is not supported by the built-in evaluator: the lookup values are "
                    "not sorted in the order this search needs (binary search on unsorted data is not implemented). "
                    "Nothing was written."),
         file_is("result.xlsx", False),
     ])],
     build_recalc_unsorted)

# 23-24. errors
task("errors_div0_after_recalc", "errors",
     "Calculate the workbook, then tell me which cells have errors.",
     [step(["recalc", "book.xlsx", "--out", "calc.xlsx", "--json"], 0, [
         j_eq(["errors"], {"Data!C1": "#DIV/0!"}),
     ]),
      step(["errors", "calc.xlsx", "--json"], 1, [
         j_eq(["error_count"], 1),
         j_eq(["not_calculated_count"], 0),
         j_eq(["findings"], [{"sheet": "Data", "cell": "C1", "formula": "=1/0", "value": "#DIV/0!"}], unordered=True),
     ])],
     build_errors_div0)

task("errors_not_calculated_formula", "errors",
     "Check book.xlsx for errors and formulas that have no result yet.",
     [step(["errors", "book.xlsx", "--json"], 1, [
         j_eq(["error_count"], 0),
         j_eq(["not_calculated_count"], 1),
         j_eq(["findings"], [{"sheet": "Data", "cell": "B1", "formula": "=SUM(A1:A3)", "value": None}]),
     ])],
     build_errors_not_calculated)

# 25. validate
task("validate_valid_and_truncated", "validate",
     "Check whether book.xlsx is a valid workbook, then check broken.xlsx (a copy cut to 1000 bytes).",
     [step(["validate", "book.xlsx", "--json"], 0, [j_eq(["verdict"], "valid")]),
      step(["validate", "broken.xlsx", "--json"], 1, [j_eq(["verdict"], "invalid")]),
      step(["sheets", "broken.xlsx", "--json"], 1, [
         stderr_has("File is not a readable .xlsx (corrupt or truncated): broken.xlsx")]),
      ],
     build_validate)

# 26-28. check-values
task("check_values_spec_example", "check-values",
     "Check the test data on Data against the rules: Name required, Age 0 to 120, Start a real date, Code like A1.",
     [step(["check-values", "book.xlsx", "--rules", "rules.json", "--sheet", "Data", "--json"], 1, [
         j_eq(["rows_checked"], 3),
         j_eq(["violations"], [
             {"sheet": "Data", "cell": "B3", "column": "Age", "rule": "min", "value": -2},
             {"sheet": "Data", "cell": "C3", "column": "Start", "rule": "type", "value": "2026-02-30"},
             {"sheet": "Data", "cell": "A4", "column": "Name", "rule": "required", "value": None},
         ], unordered=True),
     ])],
     build_check_spec)

task("check_values_unique_allowed_maxlength", "check-values",
     "Check the order list: IDs must be unique, Status must be open or closed, Qty at most 10, Note at most 5 characters.",
     [step(["check-values", "book.xlsx", "--rules", "rules.json", "--json"], 1, [
         j_eq(["rows_checked"], 3),
         j_eq(["violations"], [
             {"sheet": "Data", "cell": "A3", "column": "ID", "rule": "unique", "value": "X1"},
             {"sheet": "Data", "cell": "B3", "column": "Status", "rule": "allowed", "value": "pending"},
             {"sheet": "Data", "cell": "C3", "column": "Qty", "rule": "max", "value": 11},
             {"sheet": "Data", "cell": "D3", "column": "Note", "rule": "max_length", "value": "toolongtext"},
         ], unordered=True),
     ])],
     build_check_unique)

task("check_values_missing_column_refused", "check-values",
     "Check the data against rules that include an Email column, which is not in the sheet.",
     [step(["check-values", "book.xlsx", "--rules", "rules.json", "--json"], 1, [
         stderr_has('Column "Email" named in rules was not found in row 1.'),
     ])],
     build_check_missing_column)

# 29-30. refusals: protected file and drawing guard
task("refuse_password_protected", "refusal",
     "Show the sheets of locked.xlsx, which is password protected.",
     [step(["sheets", "locked.xlsx", "--json"], 2, [
         stderr_has("This workbook is password-protected. The skill does not open protected files."),
     ])],
     build_password)

task("refuse_drawing_guard_then_allow", "refusal",
     "Write a formula into chart.xlsx, which has a bar chart. Then do it again with the drawing rewrite allowed.",
     [step(["write", "chart.xlsx", "--patch", "patch.json", "--out", "result.xlsx", "--json"], 2, [
         stderr_has("This workbook has charts or drawing shapes that may be altered on save. "
                    "Pass --allow-drawing-rewrite to continue."),
         file_is("result.xlsx", False),
     ]),
      step(["write", "chart.xlsx", "--patch", "patch.json", "--out", "result2.xlsx",
            "--allow-drawing-rewrite", "--json"], 0, [
         j_truthy(["drawings_rewritten"]),
         cell("result2.xlsx", "Data", "D1", "=A1*2", mode="formula"),
     ])],
     build_chart)


def main():
    ids = [t["id"] for t in TASKS]
    if len(ids) != len(set(ids)):
        raise SystemExit("duplicate task ids")
    if not 25 <= len(TASKS) <= 30:
        raise SystemExit(f"expected 25 to 30 tasks, have {len(TASKS)}")
    if CASES.exists():
        shutil.rmtree(CASES)
    CASES.mkdir()
    out = []
    for t in TASKS:
        case = CASES / t["id"]
        case.mkdir()
        runnable, reason = True, ""
        try:
            t["build"](case)
        except Exception as exc:  # setup failure: task is listed as not runnable
            runnable, reason = False, f"{type(exc).__name__}: {exc}"
        out.append({"id": t["id"], "category": t["category"], "request": t["request"],
                    "runnable": runnable, "not_runnable_reason": reason, "steps": t["steps"]})
    write_json(TASKS_JSON, out)
    runnable_count = sum(1 for t in out if t["runnable"])
    print(f"wrote {len(out)} tasks to {TASKS_JSON}; runnable {runnable_count}")
    for t in out:
        if not t["runnable"]:
            print(f"  not runnable: {t['id']}: {t['not_runnable_reason']}")


if __name__ == "__main__":
    main()
